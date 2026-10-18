"""CI checks for reproducible backtest/report outputs."""
from pathlib import Path
import json
import zipfile
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from docx import Document
from regression_baseline import verify_regression_baseline

BASE = Path(__file__).resolve().parent
OUT = BASE / "output"

required = [
    BASE / "TX_flow_trend_workbook.xlsx",
    BASE / "TX_flow_trend_report.docx",
    BASE / "TX_flow_trend_report.pdf",
    OUT / "validation.json",
    OUT / "final_verification.json",
    OUT / "metrics_comparison.csv",
    OUT / "equity_comparison.png",
    OUT / "equity_comparison_log.png",
    OUT / "underwater_comparison.png",
    OUT / "indicator_timeseries.png",
    OUT / "event_windows.csv",
]
missing = [str(p.relative_to(BASE)) for p in required if not p.is_file() or p.stat().st_size == 0]
assert not missing, f"missing or empty required outputs: {missing}"

for path in [BASE / "TX_flow_trend_workbook.xlsx", BASE / "TX_flow_trend_report.docx"]:
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None, f"corrupt ZIP container: {path.name}"

wb = load_workbook(BASE / "TX_flow_trend_workbook.xlsx", read_only=True)
expected_sheets = {"資金回撤疊圖"}
assert len(wb.sheetnames) == 21, f"unexpected workbook sheet count: {len(wb.sheetnames)}"
assert {"法人濾網增量歸因", "價格單獨波動配平", "同曝險大盤比較摘要", "同曝險大盤比較逐日", "特殊期間_事後描述"}.issubset(wb.sheetnames)
assert expected_sheets.issubset(wb.sheetnames), f"missing workbook sheets: {expected_sheets - set(wb.sheetnames)}"

doc = Document(BASE / "TX_flow_trend_report.docx")
doc_text = "\n".join(p.text for p in doc.paragraphs)
assert "風險配平比較（僅列於表格）" in doc_text, "DOCX missing risk-matched table"
assert "0.293口" in doc_text and "13.67%" in doc_text, "DOCX risk-matched values differ from standard"
assert "等曝險代表路徑" not in doc_text, "DOCX retains removed comparison"
# 作業規格：書面報告需含 壹 資料介紹 ～ 伍 策略極限與檢討，且 PDF 限 6～8 頁。
for heading in ["壹、資料介紹", "貳、策略假設與訊號建構", "參、回測設定與摩擦成本", "肆、績效分析", "伍、策略極限與檢討", "特殊期間"]:
    assert heading in doc_text, f"DOCX missing required section: {heading}"

pdf_bytes = (BASE / "TX_flow_trend_report.pdf").read_bytes()
assert pdf_bytes[:5] == b"%PDF-", "report PDF signature invalid"
import re
pdf_pages = len(re.findall(rb"/Type\s*/Page(?![s/])", pdf_bytes))
assert 6 <= pdf_pages <= 8, f"report PDF must be 6-8 pages, found {pdf_pages}"
for path in [OUT / "equity_comparison.png", OUT / "equity_comparison_log.png", OUT / "underwater_comparison.png", OUT / "indicator_timeseries.png"]:
    assert path.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n", f"PNG signature invalid: {path.name}"

validation = json.loads((OUT / "validation.json").read_text(encoding="utf-8"))
final = json.loads((OUT / "final_verification.json").read_text(encoding="utf-8"))
assert validation["all_strategy_and_benchmark_pnl_reconciled"] is True
assert validation["signal_dates_before_execution"] is True
assert final["strategy"]["days"] == final["benchmark"]["days"] == validation["days"]

metrics = pd.read_csv(OUT / "metrics_comparison.csv")
strategy = metrics.iloc[0]
benchmark = metrics.iloc[1]
assert np.isclose(strategy["net_pnl"], 3277536.576, atol=1e-6)
assert np.isclose(strategy["Sharpe"], 1.3674407230136694, atol=1e-12, rtol=0)
assert np.isclose(benchmark["net_pnl"], 8268352.664, atol=1e-6)
assert np.isclose(validation["parameters"]["capital"], 2195000.0, atol=1e-6)
assert int(strategy["exposure_days"]) == 658
verify_regression_baseline({
    "period": validation["period"],
    "parameters": validation["parameters"],
    "input_sha256": validation["input_sha256"],
    "metrics": {
        "strategy.days": int(final["strategy"]["days"]),
        "strategy.net_pnl": float(strategy["net_pnl"]),
        "strategy.Sharpe": float(strategy["Sharpe"]),
        "benchmark.days": int(final["benchmark"]["days"]),
        "benchmark.net_pnl": float(benchmark["net_pnl"]),
    },
})
assert np.isclose(final["strategy"]["net_pnl"], strategy["net_pnl"], atol=1e-6, rtol=0)
assert np.isclose(final["benchmark"]["net_pnl"], benchmark["net_pnl"], atol=1e-6, rtol=0)
assert f"策略{int(strategy['trades'])}筆完整方向交易" in doc_text, "DOCX trade count differs from CSV"
source = json.loads((OUT / 'institutional_source_audit.json').read_text(encoding='utf-8'))
roll = json.loads((OUT / 'roll_gap_summary.json').read_text(encoding='utf-8'))
assert source['differing_rows'] == 2 and source['changed_signal_days'] == 0
assert roll['changed_signal_days'] == 14
assert f"{roll['changed_signal_days']}天訊號不同" in doc_text
assert np.isclose(pd.read_csv(OUT / 'flow_incremental_attribution.csv').net_difference.sum(),
                  strategy.net_pnl - pd.read_csv(OUT / 'ablation.csv').set_index('code').loc['price_only', 'net_pnl'], atol=1e-6)
ablation = pd.read_csv(OUT / "ablation.csv").set_index("code")
assert ablation.loc["flow_only", "rule"] == "交易強度MA20>0；隔日開盤多單/空手"
assert ablation.loc["price_only", "rule"] == "近月收盤>價格MA20；隔日開盤多單/空手"
events = pd.read_csv(OUT / "event_windows.csv")
strategy_daily = pd.read_csv(OUT / "strategy_daily.csv", parse_dates=["date"]).set_index("date")
assert len(events) == 8 and events.event.is_unique
for row in events.itertuples():
    window = strategy_daily.loc[row.window_start:row.window_end]
    assert len(window) == row.trading_days
    assert np.isclose(window.pnl.sum(), row.strategy_net_pnl, atol=1e-6)
    assert int(window.loc[row.worst_open_gap_date, "position_before"]) == row.held_on_worst_gap_day
market = pd.read_csv(OUT / 'market_normalized_summary.csv')
assert len(market) == 3
assert market.reference_capital.nunique() == 1
assert np.isclose(market.loc[0, 'MDD_pct'], strategy['MDD_pct'])
assert market.loc[0, 'cumulative_return'] < market.loc[2, 'cumulative_return']
assert np.isclose(market.loc[0, 'net_pnl_equivalent'], strategy['net_pnl'], atol=1e-6)
assert np.isclose(market.loc[1, 'net_pnl_equivalent'], benchmark['net_pnl'], atol=1e-6)

print(
    "CI verification passed:",
    f"strategy Sharpe={strategy['Sharpe']:.6f},",
    f"benchmark MDD={benchmark['MDD_pct']:.6f}",
)
