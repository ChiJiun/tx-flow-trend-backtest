"""從回測底稿產生八頁繁體中文PDF、Word、圖表與工作簿附頁。"""
from pathlib import Path
import pandas as pd, numpy as np, html, json, io, zipfile
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,Image
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from fontTools.ttLib import TTFont as FontToolsFont
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage

B=Path(__file__).resolve().parent;O=B/'output';read=lambda name:pd.read_csv(O/(name+'.csv'))
m=read('metrics_comparison');ms,mb=m.iloc[0],m.iloc[1];s=read('strategy_daily');bh=read('benchmark_daily')
market=read('market_normalized_daily');market_summary=read('market_normalized_summary')
mn,mt,mi=[market_summary.iloc[i] for i in range(3)]
reference_capital=float(mn.reference_capital)

def match_benchmark_volatility(benchmark_daily,target_volatility):
    """完整樣本事後波動配平；僅供報告診斷，不是可交易的小數口部位。"""
    def scaled(scale):
        pnl=benchmark_daily.pnl.astype(float)*scale
        equity=reference_capital+pnl.cumsum()
        if (equity<=0).any():raise ValueError('risk-matched equity must remain positive')
        daily_return=pnl/equity.shift(fill_value=reference_capital)
        return equity,daily_return
    lo,hi=0.0,1.0
    for _ in range(100):
        mid=(lo+hi)/2
        _,daily_return=scaled(mid)
        if daily_return.std(ddof=1)*np.sqrt(252)<target_volatility:lo=mid
        else:hi=mid
    scale=(lo+hi)/2
    equity,daily_return=scaled(scale)
    peak=equity.cummax().clip(lower=reference_capital)
    mdd=float(1-(equity/peak).min())
    years=((benchmark_daily.date.iloc[-1]-benchmark_daily.date.iloc[0]).days+1)/365.25
    cagr=float((equity.iloc[-1]/reference_capital)**(1/years)-1)
    volatility=float(daily_return.std(ddof=1)*np.sqrt(252))
    return dict(scale=scale,annual_volatility=volatility,CAGR=cagr,
                Sharpe=float(daily_return.mean()/daily_return.std(ddof=1)*np.sqrt(252)),
                MDD_pct=mdd,Calmar=cagr/mdd)
tr=read('strategy_trades');annual=read('annual_results');sens=read('cost_sensitivity');ab=read('ablation');splits=read('descriptive_time_splits')
ev=read('strategy_events');bev=read('benchmark_events');v=json.loads((O/'validation.json').read_text())
near_prices=pd.read_csv(B/'input/prices.csv',parse_dates=['date'])
contracts=pd.read_csv(B/'input/all_contracts.csv')
institutional=pd.read_csv(B/'input/institutional.csv')
candidate_count=len(pd.read_csv(B/'source_audits/all_18_exploratory_candidates.csv'))
source_audit=json.loads((O/'institutional_source_audit.json').read_text(encoding='utf-8'))
roll_audit=json.loads((O/'roll_gap_summary.json').read_text(encoding='utf-8'))
flow_attribution=read('flow_incremental_attribution')
price_start,price_end=near_prices.date.min(),near_prices.date.max()
chip_dates=pd.to_datetime(institutional['日期'])
assumptions=v['parameters']
for d in [s,bh,market]:d['date']=pd.to_datetime(d.date)
rm=match_benchmark_volatility(bh,float(ms.annual_volatility))
assert np.isclose(rm['annual_volatility'],ms.annual_volatility,atol=1e-12)
assert np.isclose(rm['scale'],0.29260755630855306,atol=1e-10)
price_vm=read('price_only_volatility_matched_summary').iloc[0]
tr['entry_date']=pd.to_datetime(tr.entry_date);tr['exit_date']=pd.to_datetime(tr.exit_date)
duration=(tr.exit_date-tr.entry_date).dt.days
pct=lambda x:f'{x*100:.2f}%';money=lambda x:f'{x:,.2f}';ratio=lambda x:f'{x:.3f}'
font=FontProperties(fname=str(B/'chinese_font.ttf'));plt.rcParams['axes.unicode_minus']=False

def validate_equity_for_log(*frames):
    for d in frames:
        if d.empty:raise ValueError('log equity chart requires non-empty data')
        equity=pd.to_numeric(d.equity,errors='coerce')
        if not np.isfinite(equity).all() or (equity<=0).any():raise ValueError('log equity chart requires strictly positive finite equity')

for water in [False,True]:
    fig,ax=plt.subplots(figsize=(9.4,3.75))
    for d,label,col,ls in [(s,'法人交易強度＋價格趨勢','#146f91','-'),(bh,'Buy-and-Hold（固定一口）','#ce7c2e','-')]:
        ax.plot(d.date,d.drawdown_pct*100 if water else d.equity/1e6,label=label,color=col,lw=1.25,ls=ls)
    ax.set_xlim(s.date.min(),s.date.max());ax.set_ylabel('回撤（%）' if water else '帳戶資金（百萬元）',fontproperties=font)
    ax.set_title('成本後回撤深度比較' if water else '成本後資金曲線：一口TX起始參考資金',fontproperties=font)
    if not water:ax.axhline(reference_capital/1e6,color='gray',ls=':',lw=.8)
    ax.legend(prop=font,frameon=False);ax.grid(alpha=.18);ax.spines[['top','right']].set_visible(False)
    fig.tight_layout();fig.savefig(O/('underwater_comparison.png' if water else 'equity_comparison.png'),dpi=200);plt.close(fig)
validate_equity_for_log(s,bh)
fig,ax=plt.subplots(figsize=(9.4,3.75))
for d,label,col,ls in [(s,'法人交易強度＋價格趨勢','#146f91','-'),(bh,'Buy-and-Hold（固定一口）','#ce7c2e','-')]:
    ax.plot(d.date,d.equity/1e6,label=label,color=col,lw=1.25,ls=ls)
ax.set_yscale('log');ax.set_xlim(s.date.min(),s.date.max());ax.set_ylabel('帳戶資金（百萬元，log scale）',fontproperties=font)
ax.set_title('成本後資金曲線（對數Y軸）：一口TX起始參考資金',fontproperties=font)
ax.legend(prop=font,frameon=False);ax.grid(alpha=.18,which='both');ax.spines[['top','right']].set_visible(False)
fig.tight_layout();fig.savefig(O/'equity_comparison_log.png',dpi=200);plt.close(fig)
# 原始盤後指標、標準化訊號與期貨價格的時間序列；陰影為策略日末持有多單的交易日。
feat_all=read('indicator_features');feat_all['date']=pd.to_datetime(feat_all.date)
first_q20=feat_all.loc[feat_all.intensity_ma20.notna(),'date'].min()
feat=feat_all[feat_all.date.between(s.date.min(),s.date.max())].reset_index(drop=True)
assert len(feat)==len(s) and feat.date.equals(s.date)
small_font=FontProperties(fname=str(B/'chinese_font.ttf'),size=7.5)
fig,(ax1,ax2)=plt.subplots(2,1,figsize=(9.4,3.75),sharex=True,gridspec_kw=dict(height_ratios=[1.25,1]))
held=(s.position_after.to_numpy()!=0)
for ax in (ax1,ax2):ax.fill_between(s.date,0,1,where=held,transform=ax.get_xaxis_transform(),color='#146f91',alpha=.13,lw=0)
ax1.plot(feat.date,feat.close,color='#333333',lw=.85,label='近月收盤（未調整）');ax1.plot(feat.date,feat.price_ma20,color='#ce7c2e',lw=.85,label='MA20');ax1.fill_between([],[],[],color='#146f91',alpha=.13,label='策略日末持有多單')
ax1.set_ylabel('TX近月（點）',fontproperties=small_font);ax1.legend(prop=small_font,frameon=False,ncol=3,loc='upper left')
ax2.bar(feat.date,feat.trade_intensity*100,width=1.0,color='#b9c9d1',label='每日交易強度Q');ax2.plot(feat.date,feat.intensity_ma20*100,color='#146f91',lw=1.0,label='20日平均Q20');ax2.axhline(0,color='gray',lw=.7)
ax2.set_ylabel('交易強度（%）',fontproperties=small_font);ax2.legend(prop=small_font,frameon=False,ncol=2,loc='upper left')
ax1.set_title('原始盤後指標與價格：三大法人交易強度Q、Q20與TX近月收盤',fontproperties=font)
for ax in (ax1,ax2):ax.grid(alpha=.18);ax.spines[['top','right']].set_visible(False);ax.tick_params(labelsize=7.5)
ax2.set_xlim(s.date.min(),s.date.max());fig.tight_layout();fig.savefig(O/'indicator_timeseries.png',dpi=200);plt.close(fig)
ew=read('event_windows')

pdfmetrics.registerFont(TTFont('TC',str(B/'chinese_font.ttf')))
styles={
 'body':ParagraphStyle('body',fontName='TC',fontSize=10.5,leading=16.6,spaceAfter=9),
 'small':ParagraphStyle('small',fontName='TC',fontSize=8.7,leading=12.4,spaceAfter=7),
 'cell':ParagraphStyle('cell',fontName='TC',fontSize=9.2,leading=13.3),
 'cellsmall':ParagraphStyle('cellsmall',fontName='TC',fontSize=8.2,leading=11.4),
 'h0':ParagraphStyle('h0',fontName='TC',fontSize=14.5,leading=20,spaceAfter=4,textColor=colors.HexColor('#0f3f57')),
 'title':ParagraphStyle('title',fontName='TC',fontSize=18,leading=24,spaceAfter=15,textColor=colors.HexColor('#145879')),
 'sub':ParagraphStyle('sub',fontName='TC',fontSize=12.2,leading=18,spaceAfter=7,textColor=colors.HexColor('#145879'))}
story=[];pages=[];current=[];texts=[]
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(1.2);sec.left_margin=sec.right_margin=Cm(1.4)
sty=doc.styles['Normal'];sty.font.name='Noto Sans TC';sty.font.size=Pt(9);sty.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans TC')
sty.paragraph_format.space_after=Pt(2);sty.paragraph_format.line_spacing=1.0
doc_page_break=False
def text(t,style='body'):
    texts.append(t);story.append(Paragraph(html.escape(t).replace('\n','<br/>'),styles[style]));current.append(t);doc.add_paragraph(t)
def title(t):
    global doc_page_break
    texts.append(t);story.append(Paragraph(html.escape(t),styles['title']));current.append('# '+t)
    paragraph=doc.add_heading(t,1)
    if doc_page_break:paragraph.paragraph_format.page_break_before=True;doc_page_break=False
def report_title(t):
    texts.append(t);story.append(Paragraph(html.escape(t),styles['h0']));current.append('# '+t);doc.add_heading(t,0)
def sub(t):text(t,'sub')
def page():
    global doc_page_break
    pages.append(current.copy());current.clear();story.append(PageBreak());doc_page_break=True
def table(rows,widths=None,style='cell'):
    for row in rows:texts.extend(str(x) for x in row)
    data=[[Paragraph(html.escape(str(x)),styles[style]) for x in row] for row in rows]
    t=Table(data,colWidths=widths or [140,365],repeatRows=1);t.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eaf2f6')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#c4d3da')),
        ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),
        ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
    story.extend([t,Spacer(1,9)]);current.extend([' | '.join(map(str,r)) for r in rows])
    tab=doc.add_table(rows=0,cols=len(rows[0]));tab.style='Table Grid'
    for row in rows:
        cells=tab.add_row().cells
        for cell,val in zip(cells,row):cell.text=str(val)
def image(name):story.append(Image(str(O/name),width=505,height=202));doc.add_picture(str(O/name),width=Cm(17))
def image_pair(left,right):
    story.append(Table([[Image(str(O/left),width=245,height=98),Image(str(O/right),width=245,height=98)]],colWidths=[252,252]))
    tab=doc.add_table(rows=1,cols=2)
    for cell,name in zip(tab.rows[0].cells,[left,right]):cell.paragraphs[0].add_run().add_picture(str(O/name),width=Cm(8.1))

report_title('臺灣期交所盤後籌碼量化分析與交易策略回測：三大法人交易強度20日＋價格趨勢（TX近月一口多單／空手）')
title('壹、資料介紹：來源、欄位定義與經濟意涵')
text('研究假說：三大法人持續的淨買進活動，若同時伴隨價格上升趨勢，可能比單日部位異常更能辨識可持續的多頭環境。本研究以20日法人交易強度與20日價格均線建立一口多單／空手規則，以加權股價指數作大盤價格參考，另與計入轉倉成本的台指期近月Buy-and-Hold比較。')
table([['資料','來源、期間與清理'],['TX日盤行情',f'TAIFEX「期貨每日交易行情」年度ZIP及月檔；{price_start:%Y/%m/%d}～{price_end:%Y/%m/%d}，{near_prices.date.nunique():,}日、{len(contracts):,}筆實際月契約、{contracts.contract.nunique()}個契約。僅保留一般交易時段、六碼月契約。'],['三大法人交易',f'TAIFEX「三大法人－區分各期貨契約」格式原始資料，並以FinMind TX歷史資料逐筆比對；{chip_dates.min():%Y/%m/%d}～{chip_dates.max():%Y/%m/%d}，{chip_dates.nunique():,}日、{len(institutional):,}筆三類法人資料。'],['正式回測期間',f'{ms.start}～{ms.end}，共{int(ms.days):,}日。起始日沿用先前候選比較的共同期間（最長需60日暖身）；本策略只需20日，Q20自{first_q20:%Y/%m/%d}起即可計算。'],['缺漏處理',f'行情與法人資料起點不同；法人資料前的{near_prices.loc[near_prices.date<chip_dates.min(),"date"].nunique():,}個行情交易日未補造籌碼。較早官方交易金額淨額空白，但本策略只用完整的交易口數。']])
text(f'較新期間以TAIFEX格式原始資料優先，以降低不同資料來源歷史值差異造成的影響；三類法人日期均與行情交易日曆一致。資料未混入MTX、週契約、夜盤或價差交易。原始月契約共有{contracts[['open','high','low','close']].isna().any(axis=1).sum()}筆OHLC缺值紀錄，但策略及基準實際使用的進出、轉倉與評價價格均完整。','small')
sub('期交所三大法人期貨交易資訊：原始欄位字典與公布時點')
table([['原始欄位','定義、單位與市場意涵'],['日期／商品名稱／身份別','交易日；臺股期貨；外資及陸資、投信、自營商。三類均為眾多機構的合計，淨額只是互抵結果，不代表單一機構或整類法人一致看多或看空。[2]'],['多方交易口數／空方交易口數','當日買進成交口數／賣出成交口數（口），屬成交「流量」。本研究核心欄位，用以計算淨交易口數B與總交易口數V。'],['多空交易口數淨額','多方減空方交易口數（口）；正值為當日淨買進。已核對全數等於兩欄相減。'],['多方／空方／淨額交易契約金額（千元）','成交口數×成交價×乘數；早期官方淨額欄位空白，本研究未使用金額欄位。'],['多方／空方未平倉口數、多空未平倉口數淨額','收盤後仍未沖銷的留倉口數（口），屬部位「存量」，常被解讀為法人方向性曝險。本研究保留於原始資料供對照，訊號只用當日成交流量，以避免存量受避險與套利部位長期累積影響。'],['公布時點／最早可用時點','期交所於交易日收盤後（約下午3時）公布T日資料；本研究保守採T+1日日盤08:45開盤為最早執行時點，不假設可於T日夜盤成交。']])
text(f'本策略由{candidate_count}個候選規則中篩選後固定參數重新回測，因此結果屬探索性全樣本分析，尚未以完全未參與策略選擇的資料進行獨立樣本外驗證。此限制在後續結論中另行討論。','small')
text(f'以首日一口TX名目價值{money(reference_capital)}元作共同起始資金，策略成本後累積報酬{pct(ms.cumulative_return)}、MDD {pct(ms.MDD_pct)}；同期加權指數價格報酬{pct(mi.cumulative_return)}、MDD {pct(mi.MDD_pct)}。固定一口TX買進持有為可交易期貨對照。各曲線逐日累加損益，回撤從各自的歷史高點計算。','small')
page();title('壹（續）指標定義、經濟意涵與資料時間序列')
table([['欄位／指標','定義與用途'],['日期／實際契約月份','交易日、TXYYYYMM；每日開、高、低、收都取同一實際月契約，未回溯調整。'],['多方／空方交易口數','法人當日買進／賣出成交口數。與仍未沖銷的未平倉口數是不同資料。[2]'],['法人淨交易口數 B','外資及陸資、投信、自營商的多方交易口數減空方交易口數，再合計。'],['法人總交易口數 V','三類法人多方與空方交易口數合計。只代表這三類法人，不是全市場成交量。'],['交易強度 Q=B/V','以口數標準化淨買進活動；B可為負，V為正，Q介於−1與1之間。各類法人的成交流量可能含方向、避險與套利。'],['Q20／MA20','Q20為當日及前19交易日Q的算術平均；MA20為當日及前19日近月收盤均價。']])
text('例如三類法人合計淨買進5,000口、買賣總口數200,000口，交易強度為2.5%。20日平均為正，表示近期平均成交傾向淨買方；但它不是新增多單口數，也不能把多次買賣直接視為留倉增加。')
text('經濟直覺：持續淨買進可能反映資金配置或風險承擔；20日平均降低單日成交雜訊，價格高於MA20則要求市場價格同時呈現偏強狀態。這是「法人買方活動與價格趨勢一致」的條件假說，不預設它必然有預測能力。')
image('indicator_timeseries.png')
text(f'圖一：上圖為TX近月收盤（未調整）與MA20，下圖為每日交易強度Q（長條）與Q20（實線）；陰影為策略日末持有多單的交易日（共{int(ms.exposure_days)}日，占{pct(ms.exposure_days/ms.days)}）。Q的單日雜訊大（中位絕對值{feat.trade_intensity.abs().median()*100:.2f}%），Q20全期介於{feat.intensity_ma20.min()*100:+.2f}%與{feat.intensity_ma20.max()*100:+.2f}%之間、{pct((feat.intensity_ma20>0).mean())}的交易日為正，因此零門檻實際上是「近期法人成交偏向淨買方」的方向判斷，而非強度分級。','small')
page();title('貳、策略假設與訊號建構')
text('第t日盤後：若Q20_t>0且收盤C_t>MA20_t，則目標多一口；其他情況目標空手。Q20=0或C=MA20均不持有。訊號只在t+1交易日08:45開盤執行，不使用t+1收盤決定t+1開盤交易；T日盤後資料不假設可於T日夜盤成交。')
text('原本空手且目標多單：買進；原本多單且條件持續：跨日持有、不加碼；任一條件消失：次日開盤全部平倉。無放空、無日內停損停利。OHLC保留作稽核，但不能重建夜盤及盤中成交路徑，所以不加入未經驗證的盤中觸價出場。')
text('交易日配對按官方行情日曆，不以自然日加一天。條件產生、成交、持倉及轉倉均可在逐日底稿追溯；可能失效情境：盤整反覆穿越均線、急跌後V形反彈，以及法人成交流量由避險或套利主導的期間（見伍）。')
title('參、回測設定與摩擦成本')
text('策略按訊號隔日開盤進出；若跨月仍持有，於到期前一個交易日13:45平舊約並買次月。基準在首日開盤買一口，持續持有並採相同提前轉倉路徑。到期日期由官方契約存續底稿核對第三個星期三／休市順延規則。[1]')
text('提前一日換倉可避免到期近月13:30、次月13:45收盤時點不同；不用最後結算價替代成交價。若策略在換倉日開盤已退出，就不再換倉。兩者在樣本結束時收盤全平，這是帳戶終止清算，並非日常開盤買、收盤賣規則。')
table([['設定','策略與基準共同口徑'],['起始資金／口數／乘數',f'共同起始資金{money(reference_capital)}元，等於首日TX開盤價乘以每點200元；持倉時固定一口，不加碼。'],['手續費','每邊每口50元，假設為券商總手續費；未另重複加交易所經手費。[3]'],['滑價','每次成交2點：買進價+2、賣出價−2。初始進出、轉倉與期末平倉皆計。'],['期交稅','每邊按含滑價成交價×200×0.00002。連續金額近似，未逐筆取整。[3]'],['資金與槓桿','起始資金等於一口名目價值，名目槓桿約1倍；若改以保證金為分母，報酬與回撤比例會同比例放大。起始資金不代表歷史保證金要求，未重建逐日保證金與追繳；全期策略與基準權益均維持正值。'],['日末評價與報酬','以共同起始資金逐日累加扣除全部成本後的淨損益，得到日末權益曲線；回撤取權益相對歷史高點的最大跌幅。日報酬為當日淨損益／前日權益，含空手日。'],['年化與利息','現金利息及rf=0；Sharpe採252日；CAGR依實際起訖含首尾日／365.25。']])
text('持有日損益按同一實際契約今日收盤減昨日評價價；退出日先計舊持倉至開盤的損益，進場日從開盤計至收盤。換倉日先結束舊約損益，再將新約收盤設為新評價基準，成交各扣成本。新舊約價差不直接列為本日損益，也不另外重複扣費。')
text(f'策略147筆完整方向交易、37次轉倉、368個成交腿；基準98段契約交易、97次轉倉、196個成交腿。策略總成本{money(ms.cost_total)}元，基準{money(mb.cost_total)}元。逐日損益、方向交易及實際契約分段損益皆獨立核對一致。')
page();title('肆、績效分析：量化指標與年度比較')
rows=[['成本後指標','法人交易強度＋趨勢','Buy-and-Hold基準']]
for label,key,fmt in [('累積總報酬','cumulative_return',pct),('CAGR','CAGR',pct),('期末資金（元）','final_equity',money),('交易次數／段數','trades',lambda z:str(int(z))),('勝率','win_rate',pct),('Profit Factor','profit_factor',ratio),('最大回撤（元）','MDD_amount',money),('最大回撤（%）','MDD_pct',pct),('最長回撤：交易日','longest_drawdown_trading_days',lambda z:str(int(z))),('最長回撤：日曆日','longest_drawdown_calendar_days',lambda z:str(int(z))),('Sharpe（rf=0）','Sharpe',ratio),('Calmar','Calmar',ratio),('年化波動','annual_volatility',pct),('全部成本（元）','cost_total',money)]:rows.append([label,fmt(ms[key]),fmt(mb[key])])
table(rows,[175,165,165])
text('同期間、同一口、同起始資金。固定一口Buy-and-Hold是可直接執行的曝險比較；策略有658個日末持倉日，基準幾乎全期持有，兩者承擔市場風險的時間不同。MDD金額及比例各自取全期最大值，可能發生於不同日期。','small')
ar=[['年度','策略損益（元）','策略報酬','基準損益（元）','基準報酬']]
for year,g in annual.groupby('year'):
    q=g[g.series=='strategy'].iloc[0];r=g[g.series=='benchmark'].iloc[0]
    ar.append([int(year),f'{q.net_pnl:,.0f}',pct(q.annual_return),f'{r.net_pnl:,.0f}',pct(r.annual_return)])
table(ar,[45,140,90,140,90])
text('2018、2026為部分年度；年度報酬未年化，以該年首日損益前權益為分母。PF=淨正損益合計／淨負損益絕對值合計。最長回撤包含期末未恢復區間，含初始本金高點及空手日。','small')
page();title('肆（續）資產淨值累積曲線、回撤水下圖與風險配平診斷')
image_pair('equity_comparison.png','equity_comparison_log.png')
text('圖二（Equity Curve）：上方兩圖只比較策略與固定一口 Buy-and-Hold；左為線性Y軸、右為對數Y軸。風險配平 benchmark 不加入圖表，避免與可直接交易的固定一口路徑混淆。','small')
image('underwater_comparison.png')
text(f'圖三（Underwater Chart）：回撤仍只比較策略與固定一口基準。策略MDD為{pct(ms.MDD_pct)}，固定一口基準為{pct(mb.MDD_pct)}。','small')
sub('風險配平比較（僅列於表格）')
table([['指標','固定一口基準','風險配平基準','V04策略'],['部位','1.000口',f'{rm["scale"]:.3f}口','持倉時1口'],['年化波動',pct(mb.annual_volatility),pct(rm['annual_volatility']),pct(ms.annual_volatility)],['CAGR',pct(mb.CAGR),pct(rm['CAGR']),pct(ms.CAGR)],['Sharpe',ratio(mb.Sharpe),ratio(rm['Sharpe']),ratio(ms.Sharpe)],['MDD',pct(mb.MDD_pct),pct(rm['MDD_pct']),pct(ms.MDD_pct)],['Calmar',ratio(mb.Calmar),ratio(rm['Calmar']),ratio(ms.Calmar)]],[105,135,135,130])
text(f'風險配平基準將固定一口 benchmark 的每日毛損益與全部成本同比例縮放至約{rm["scale"]:.3f}口，使年化波動與策略同為{pct(ms.annual_volatility)}。配平後 Sharpe 為{ratio(rm["Sharpe"])}，策略為{ratio(ms.Sharpe)}。此係數由完整樣本事後估計，且TX不能交易小數口，因此只用於同風險尺度比較。','small')
page();title('肆（續）成本敏感度、條件診斷與特殊期間')
sub('滑價增加後是否仍保留成果')
rows=[['每邊滑價','策略報酬','策略Sharpe','基準報酬','基準Sharpe']]
for slip,g in sens.groupby('slippage_per_leg'):
    q=g[g.series=='strategy'].iloc[0];r=g[g.series=='benchmark'].iloc[0]
    rows.append([f'{slip:g}點',pct(q.cumulative_return),ratio(q.Sharpe),pct(r.cumulative_return),ratio(r.Sharpe)])
table(rows,[75,105,110,105,110])
high_slip=sens[(sens.series=='strategy') & (sens.slippage_per_leg==sens.slippage_per_leg.max())].iloc[0]
text(f'基準{v["parameters"]["slippage_per_leg"]:g}點設定下，策略毛損益{money(ms.gross_pnl)}元，成本{money(ms.cost_total)}元，占毛獲利{pct(ms.cost_total/ms.gross_pnl)}。其中手續費{money(ms.fee)}元、期交稅{money(ms.tax)}元、滑價{money(ms.slippage_cost)}元。{high_slip.slippage_per_leg:g}點滑價時Sharpe為{ratio(high_slip.Sharpe)}；固定滑價不能完整模擬極端行情的成交失敗與流動性惡化。')
sub('拆開兩個條件：事後消融，不是樣本外')
rows=[['規則','報酬','Sharpe','MDD'],['交易強度＋價格趨勢',pct(ms.cumulative_return),ratio(ms.Sharpe),pct(ms.MDD_pct)]]
for r in ab.itertuples():rows.append([r.series,pct(r.cumulative_return),ratio(r.Sharpe),pct(r.MDD_pct)])
table(rows,[220,95,95,95])
text('消融表列出雙條件、單獨法人條件及單獨價格條件在同一期間的結果；報酬、回撤與曝險可分別對照。消融是在選定策略後作的描述，不能消除先前候選挑選造成的偏誤。')
sub('交易分布與成果集中')
strategy_annual=annual[annual.series=='strategy'].sort_values('year')
recent_years=strategy_annual.year.tail(2).astype(int).tolist()
recent_pnl=strategy_annual[strategy_annual.year.isin(recent_years)].net_pnl.sum()
parts=[]
for r in splits[splits.series=='strategy'].itertuples():
    parts.append(f'{r.period}淨利{money(r.net_pnl)}元、描述性Sharpe {ratio(r.Sharpe)}')
negative_years=strategy_annual.loc[strategy_annual.net_pnl<0,'year'].astype(int).tolist()
negative_years_text='、'.join(map(str,negative_years)) if negative_years else '無'
text(f'{int(ms.trades)}筆交易勝率{pct(ms.win_rate)}、PF {ratio(ms.profit_factor)}；持有日曆日平均{duration.mean():.2f}日、中位數{duration.median():.0f}日，最長{duration.max()}日。前五筆獲利合計占總淨利{pct(tr.nlargest(5,"pnl").pnl.sum()/ms.net_pnl)}；最近兩個曆年（{recent_years[0]}、{recent_years[-1]}）合計占總淨利{pct(recent_pnl/ms.net_pnl)}，仍需注意少數交易及近期年份的影響。'+'；'.join(parts)+f'。年度虧損年份：{negative_years_text}。')
sub('特殊期間：急跌與跳空事件下的訊號狀態（事後描述，不用於調整規則）')
rows=[['事件視窗','交易日','起點部位','持倉日','策略損益（元）','基準損益（元）','最深開盤跳空（日期）','跳空日持倉']]
for r in ew.itertuples():rows.append([r.event,int(r.trading_days),'多單' if r.position_at_window_start else '空手',int(r.strategy_exposure_days),f'{r.strategy_net_pnl:,.0f}',f'{r.benchmark_net_pnl:,.0f}',f'{r.worst_open_gap_pct*100:.1f}%（{r.worst_open_gap_date[5:]}）','是' if r.held_on_worst_gap_day else '否'])
table(rows,[120,36,44,36,66,70,96,37],'cellsmall')
flat_on_gap=int((ew.held_on_worst_gap_day==0).sum());worst_days=s.nsmallest(3,'pnl')
worst_text='、'.join(f'{d.date:%Y/%m/%d}（{"進場日" if d.position_before==0 else "持倉日"}，{d.pnl:,.0f}元）' for d in worst_days.itertuples())
text(f'事件日期以公開重大事件為起點、由行情日曆截取；末兩段以樣本內最深單日跳空由資料辨識，均未回饋到規則。{len(ew)}個視窗中{flat_on_gap}個在最深跳空日為空手：急跌前價格多已跌破MA20或Q20轉負而退出，策略是「不在場」而非預知事件，同一機制也使其在反彈初期缺席（錯過幅度與最差單日見伍）。','small')
page();title('伍、策略極限與檢討')
sub('一、法人交易指標的解釋限制')
text('本研究的交易強度衡量三大法人的當日淨成交流量，而非收盤後仍持有的未平倉部位。以多空交易總口數作為分母，可降低不同交易活躍程度對原始淨口數的影響；然而三類法人交易同時可能包含方向投資、股票與選擇權避險、套利及造市活動，因此公開彙總資料無法識別各機構的實際交易動機。[2]')
text('20日平均用於降低單日成交流量雜訊，再以價格高於20日均線作為趨勢確認。回測中較高的Sharpe僅代表此規則在目前樣本中的歷史風險調整績效；由於策略採多單／空手架構，績效同時受到長期多頭趨勢與隔夜報酬影響，不能將全部成果歸因於法人資訊優勢。')
text(f'固定一口Buy-and-Hold的MDD {pct(mb.MDD_pct)}高於策略的{pct(ms.MDD_pct)}，但其年化波動{pct(mb.annual_volatility)}亦遠高於策略的{pct(ms.annual_volatility)}；縮放至約{rm["scale"]:.3f}口的配平基準Sharpe為{ratio(rm["Sharpe"])}、MDD為{pct(rm["MDD_pct"])}（見肆續表），這只控制風險尺度，小數口並非可直接交易的部位。')
text(f'逐筆比對兩份法人來源共{source_audit["compared_rows"]:,}筆，其中{source_audit["differing_rows"]}筆在2026/07/06不同，造成{source_audit["affected_ma20_days"]}個交易日的Q20變化，但本樣本{source_audit["changed_signal_days"]}天訊號改變。歷史逐日初次發布快照與修訂軌跡未留存，故不能聲稱全期資料時點已獨立驗證。')
text(f'近月未調整價格均線有換月價差敏感度：以前一日兩契約收盤重疊價差調整訊號後，{roll_audit["changed_signal_days"]}天訊號不同，淨利{money(roll_audit["adjusted_net_pnl"])}元、Sharpe {ratio(roll_audit["adjusted_Sharpe"])}；正式策略為{money(roll_audit["raw_net_pnl"])}元、{ratio(roll_audit["raw_Sharpe"])}。此項診斷只變更訊號，實際成交仍按原契約。')
text(f'價格單獨策略按完整樣本事後配平至相同年波動的TX-equivalent為{price_vm.position_equivalent:.3f}，淨利{money(price_vm.net_pnl)}元、Sharpe {ratio(price_vm.Sharpe)}；策略為{money(ms.net_pnl)}元、{ratio(ms.Sharpe)}。這是風險尺度與曝險差異的描述，配平小數口不能直接交易，亦非樣本外證據。')
sub('二、盤整與快速反轉環境')
text('在盤整、急漲急跌或V形反轉環境中，20日平均具有滯後性，價格亦可能在MA20附近反覆穿越。當條件短暫失效後退出、再度成立後重新進場，可能造成追價、較低價格退出或錯失快速反彈。雙條件一致仍可能漏接行情或反覆進出。')
text('本回測未加入事後最佳化的固定點數停損。現有資料僅含日盤OHLC，缺乏完整夜盤價格路徑、委託排隊及極端行情成交資訊，因此若加入盤中觸價停損，將引入無法可靠驗證的成交假設。策略僅依訊號消退、轉倉與期末清算執行出場。')
sub('三、隔夜跳空與流動性風險')
text('本策略實際持有隔夜部位，因此海外市場、匯率與突發事件可能在臺指期日盤重新開市前改變價格。即使隔日日盤開盤已產生退出需求，仍須承受前一日收盤至下一交易日開盤的跳空風險；若開盤流動性不足，實際成交損失亦可能高於固定2點滑價假設。')
rebound=ew.dropna(subset=['near_month_move_gap_close_to_reentry_open_pct']);top_rebound=rebound.loc[rebound.near_month_move_gap_close_to_reentry_open_pct.idxmax()]
text(f'特殊期間表顯示策略在{len(ew)}個急跌視窗中有{flat_on_gap}個於最深跳空日為空手，這來自雙條件在下跌初期即退出，而非對事件的預測；代價是反彈初期缺席：從最深跳空日到重新進場，中位數需{rebound.trading_days_to_reentry.median():.0f}個交易日，期間近月平均已回升{rebound.near_month_move_gap_close_to_reentry_open_pct.mean()*100:.1f}%，{top_rebound.event}更達{top_rebound.near_month_move_gap_close_to_reentry_open_pct*100:.1f}%。樣本內三個最差單日為{worst_text}，均發生在快速反轉時的進場或持倉日，跳空風險集中於重新進場後而非持倉中段。因此{pct(ms.MDD_pct)}的日末MDD不應解讀為單筆交易或盤中可能遭受的最大損失；共同起始資金並未重建歷史保證金與追繳，評估實際可行性仍需納入夜盤價格、盤中極端波動、開盤成交品質與逐日保證金需求。')
page();title('陸、結論與後續驗證')
sub('一、績效結論')
text(f'以{money(reference_capital)}元建立兩條權益曲線，逐日累加淨損益後，策略Sharpe為{ratio(ms.Sharpe)}、Calmar為{ratio(ms.Calmar)}，固定一口Buy-and-Hold為{ratio(mb.Sharpe)}與{ratio(mb.Calmar)}；策略MDD為{pct(ms.MDD_pct)}，基準為{pct(mb.MDD_pct)}。策略累積報酬{pct(ms.cumulative_return)}低於基準{pct(mb.cumulative_return)}。固定一口比較仍有持倉時間差異：策略持倉日約為基準三分之一。')
text(f'將 benchmark 事後縮放至與策略相同的{pct(ms.annual_volatility)}年化波動後，配平 benchmark 的Sharpe為{ratio(rm["Sharpe"])}、MDD為{pct(rm["MDD_pct"])}、Calmar為{ratio(rm["Calmar"])}；策略分別為{ratio(ms.Sharpe)}、{pct(ms.MDD_pct)}與{ratio(ms.Calmar)}。策略在同風險尺度下仍較佳，但配平係數使用完整樣本且為小數口，只是診斷，不是可部署的交易基準。')
text('策略最長回撤期間為482個交易日，高於基準的366日，顯示較低回撤幅度不等同於較短的資金恢復時間。勝率約51.70%、Profit Factor約2.197，績效並非依賴高勝率，而較依賴部分趨勢延續期間的較大獲利；前五筆交易與近兩年績效占比較高，亦顯示成果存在集中性。')
text(f'加權指數價格報酬為{pct(mi.cumulative_return)}、MDD {pct(mi.MDD_pct)}，但未含股息或追蹤成本，僅作市場參考。')
sub('二、樣本選擇偏誤與驗證要求')
text(f'本策略由{candidate_count}個候選規則比較後選出，並非在完全未觀察資料的情況下事先固定。因此，即使滑價敏感度與時間分組結果仍為正，也不能視為獨立樣本外證據。目前的Sharpe可能包含策略挑選造成的選擇偏誤，應在解讀績效時明確納入。')
text('後續驗證應先凍結20日交易強度、20日價格均線、零門檻及成本假設，再以未參與策略設計的新資料進行樣本外檢驗。若進一步測試結算日排除、均線參數或停損規則，應事前限定假說與參數範圍，並完整報告所有結果，以降低多重嘗試造成的偏誤。')
text('綜合而言，策略規則可追溯且歷史成本後報酬為正，但同期大盤與TX買進持有累積報酬均較高；策略持倉較少且回撤較低，未能據此認定法人訊號創造穩定超額報酬。仍須完成獨立樣本外驗證、夜盤與盤中風險及資料發布時點檢查。')
sub('三、資料來源與計算口徑')
text('[1] 臺灣期貨交易所，臺股期貨契約規格：https://www.taifex.com.tw/cht/2/tX\n[2] 臺灣期貨交易所，三大法人交易資訊：https://www.taifex.com.tw/cht/3/futContractsDateView\n[3] 臺灣期貨交易所，費率表：https://www.taifex.com.tw/cht/4/feeSchedules\n歷史行情：https://www.taifex.com.tw/cht/3/dlFutDailyMarketView\nFinMind：https://finmind.github.io/tutor/TaiwanMarket/Derivative/\n加權指數：臺灣證券交易所每日市場成交資訊與發行量加權股價指數歷史資料，來源與雜湊見source_audits/taiex_sources.json。','small')
text('Sharpe＝日報酬平均／日報酬樣本標準差×√252；Calmar＝CAGR／最大回撤比例。權益曲線以首日一口TX開盤名目價值起算，逐日累加淨損益，MDD取曲線相對歷史高點的最大跌幅；風險配平 benchmark 僅作同風險尺度診斷。','small')
pages.append(current.copy())
cmap=FontToolsFont(B/'chinese_font.ttf').getBestCmap();missing=sorted(set(ch for t in texts for ch in t if not ch.isspace() and ord(ch) not in cmap))
assert not missing,missing
def footer(can,doc):
    can.setStrokeColor(colors.HexColor('#c4d3da'));can.line(45,35,550,35);can.setFont('TC',8);can.setFillColor(colors.HexColor('#526977'))
    can.drawString(45,23,'法人交易強度＋價格趨勢｜八年共同期間與轉倉基準');can.drawRightString(550,23,f'{doc.page} / 8')
SimpleDocTemplate(str(B/'TX_flow_trend_report.pdf'),pagesize=(595.28,841.89),leftMargin=45,rightMargin=45,topMargin=40,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
doc.save(B/'TX_flow_trend_report.docx');(B/'report_text.md').write_text('\n\n---\n\n'.join('\n\n'.join(p) for p in pages),encoding='utf-8')
# 工作簿附上中文欄位說明、圖表與交易期間資訊。
w=load_workbook(B/'TX_flow_trend_workbook.xlsx')
for name in ['欄位中文說明','資金回撤疊圖']:
    if name in w.sheetnames:del w[name]
ws=w.create_sheet('欄位中文說明');ws.append(['欄位','定義'])
dictionary=[('signal_date','昨日盤後訊號日期'),('institution_net_trade_qty','三法人淨交易口數合計'),('institution_total_trade_qty','三法人買進及賣出交易口數合計'),('trade_intensity','淨交易口數合計除以總交易口數'),('intensity_ma20','含訊號日的20日交易強度平均'),('signal_close / price_ma20','訊號日近月收盤／其20日平均'),('target / buy_condition','當日目標部位／昨日訊號條件'),('position_before / position_after','昨日收盤部位／本日收盤部位'),('contract_after','日末持有實際契約；空手時空白'),('gross_pnl / pnl','毛損益／淨損益，均為新臺幣'),('fee / tax / slippage_cost / cost','手續費／稅／滑價／逐日總成本'),('total_cost','完整方向交易的全部成本'),('entry_raw / exit_raw','未含滑價的實際進出成交價'),('entry_execution / exit_execution','含滑價進出價'),('episode_id / segment_id','完整方向交易／實際契約分段編號'),('roll / reason','是否轉倉／成交原因'),('equity / daily_return','以2,195,000元起始，逐日累加淨損益形成日末權益；日報酬為當日淨損益除以前日權益'),('drawdown_pct / drawdown_amount','相對權益曲線歷史高點的回撤比例／金額')]
for row in dictionary:ws.append(row)
ws.column_dimensions['A'].width=46;ws.column_dimensions['B'].width=66;ws.freeze_panes='A2'
ws=w.create_sheet('資金回撤疊圖');ws['A1']=f'同期間、固定一口／空手；共同起始資金為首日一口TX開盤名目價值{reference_capital:,.0f}元。'
for name,cell in [('indicator_timeseries.png','A3'),('equity_comparison.png','A31'),('equity_comparison_log.png','A59'),('underwater_comparison.png','A87')]:
    im=ExcelImage(O/name);im.width=940;im.height=374;ws.add_image(im,cell)
buf=io.BytesIO();w.save(buf);payload=buf.getvalue()
with zipfile.ZipFile(io.BytesIO(payload)) as z:assert z.testzip() is None
(B/'TX_flow_trend_workbook.xlsx').write_bytes(payload)
print('報告、Word、圖表及工作簿完成；字形檢查通過。')
